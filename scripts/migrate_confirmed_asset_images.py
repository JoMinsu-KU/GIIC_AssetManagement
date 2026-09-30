import argparse
import csv
import hashlib
import json
import os
import shutil
import subprocess
import sys
import warnings
from collections import defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_WORKBOOK = PROJECT_ROOT / "outputs" / "confirmed_asset_image_mapping.xlsx"
DEFAULT_UPLOAD_DIR = PROJECT_ROOT / "uploads"
SUPPORTED_EXTENSIONS = {".jpg", ".jpeg", ".png", ".webp"}


def parse_args():
    parser = argparse.ArgumentParser(
        description="Copy confirmed asset images and register them in MySQL."
    )
    parser.add_argument("--apply", action="store_true", help="Apply file and DB changes.")
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument("--upload-dir", type=Path, default=DEFAULT_UPLOAD_DIR)
    parser.add_argument("--container", default="asset_management_db")
    parser.add_argument("--database", default="asset_management")
    return parser.parse_args()


def run_mysql(container, database, sql):
    command = [
        "docker",
        "exec",
        "-i",
        container,
        "sh",
        "-c",
        (
            'MYSQL_PWD="$MYSQL_ROOT_PASSWORD" '
            f"mysql -uroot --default-character-set=utf8mb4 -N -B {database}"
        ),
    ]
    result = subprocess.run(
        command,
        input=sql,
        text=True,
        capture_output=True,
        encoding="utf-8",
    )
    if result.returncode != 0:
        raise RuntimeError(result.stderr.strip() or "MySQL command failed.")
    return result.stdout


def sql_string(value):
    return "'" + str(value).replace("\\", "\\\\").replace("'", "''") + "'"


def file_sha256(path):
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def load_mapping(workbook_path):
    warnings.filterwarnings("ignore", message="Workbook contains no default style")
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    if "Confirmed Mapping" not in workbook.sheetnames:
        raise ValueError("The workbook does not contain a 'Confirmed Mapping' sheet.")

    rows = list(workbook["Confirmed Mapping"].iter_rows(values_only=True))
    if not rows:
        raise ValueError("The mapping sheet is empty.")

    headers = [str(value).strip() if value is not None else "" for value in rows[0]]
    required = {"asset_no", "image_path", "asset_name", "model"}
    missing_headers = required - set(headers)
    if missing_headers:
        raise ValueError(f"Missing required columns: {sorted(missing_headers)}")

    records = []
    for excel_row, row in enumerate(rows[1:], start=2):
        if not any(value is not None for value in row):
            continue
        record = dict(zip(headers, row))
        asset_no = str(record["asset_no"]).strip()
        source_path = Path(str(record["image_path"]).strip()).resolve()
        extension = source_path.suffix.lower()

        if not source_path.is_file():
            raise FileNotFoundError(f"Source image not found at row {excel_row}: {source_path}")
        if source_path.stat().st_size == 0:
            raise ValueError(f"Source image is empty at row {excel_row}: {source_path}")
        if extension not in SUPPORTED_EXTENSIONS:
            raise ValueError(f"Unsupported image extension at row {excel_row}: {extension}")

        digest = file_sha256(source_path)
        filename = f"{asset_no}_{digest[:16]}{extension}"
        records.append(
            {
                "excel_row": excel_row,
                "asset_no": asset_no,
                "asset_name": str(record.get("asset_name") or "").strip(),
                "model": str(record.get("model") or "").strip(),
                "source_path": source_path,
                "source_sha256": digest,
                "filename": filename,
                "image_url": f"/uploads/{filename}",
            }
        )

    duplicate_keys = defaultdict(list)
    for record in records:
        duplicate_keys[(record["asset_no"], record["source_sha256"])].append(record["excel_row"])
    duplicates = {key: rows for key, rows in duplicate_keys.items() if len(rows) > 1}
    if duplicates:
        raise ValueError(f"Duplicate asset/image mappings found: {duplicates}")
    return records


def load_database_state(container, database):
    asset_sql = "SELECT asset_no, id, name, COALESCE(model, '') FROM assets ORDER BY asset_no;"
    image_sql = (
        "SELECT a.asset_no, ai.image_url, CAST(ai.is_primary AS UNSIGNED) "
        "FROM asset_images ai JOIN assets a ON a.id = ai.asset_id "
        "ORDER BY a.asset_no, ai.id;"
    )

    assets = {}
    for line in run_mysql(container, database, asset_sql).splitlines():
        parts = line.split("\t")
        if len(parts) != 4:
            raise ValueError(f"Unexpected asset query row: {line}")
        assets[parts[0]] = {
            "id": int(parts[1]),
            "name": parts[2],
            "model": parts[3],
        }

    images = defaultdict(list)
    for line in run_mysql(container, database, image_sql).splitlines():
        parts = line.split("\t")
        if len(parts) != 3:
            raise ValueError(f"Unexpected image query row: {line}")
        images[parts[0]].append(
            {"image_url": parts[1], "is_primary": parts[2] == "1"}
        )
    return assets, images


def validate_database_matches(records, assets):
    failures = []
    for record in records:
        asset = assets.get(record["asset_no"])
        if not asset:
            failures.append(f"{record['asset_no']}: asset not found")
            continue
        if record["asset_name"] != asset["name"] or record["model"] != asset["model"]:
            failures.append(
                f"{record['asset_no']}: workbook=({record['asset_name']}, {record['model']}), "
                f"database=({asset['name']}, {asset['model']})"
            )
    if failures:
        sample = "\n".join(failures[:20])
        raise ValueError(f"Asset metadata validation failed:\n{sample}")


def prepare_actions(records, upload_dir, assets, existing_images):
    actions = []
    known_urls = {
        asset_no: {image["image_url"] for image in images}
        for asset_no, images in existing_images.items()
    }
    primary_exists = {
        asset_no: any(image["is_primary"] for image in images)
        for asset_no, images in existing_images.items()
    }

    for record in records:
        target_path = upload_dir / record["filename"]
        if target_path.exists() and file_sha256(target_path) != record["source_sha256"]:
            raise ValueError(f"Target file has different content: {target_path}")

        asset_no = record["asset_no"]
        already_registered = record["image_url"] in known_urls.get(asset_no, set())
        is_primary = False
        if not already_registered and not primary_exists.get(asset_no, False):
            is_primary = True
            primary_exists[asset_no] = True

        actions.append(
            {
                **record,
                "asset_id": assets[asset_no]["id"],
                "target_path": target_path,
                "copy_required": not target_path.exists(),
                "insert_required": not already_registered,
                "is_primary": is_primary,
            }
        )
    return actions


def copy_files(actions, upload_dir):
    upload_dir.mkdir(parents=True, exist_ok=True)
    copied = []
    try:
        for action in actions:
            if not action["copy_required"]:
                continue
            target_path = action["target_path"]
            temp_path = target_path.with_name(f".{target_path.name}.tmp-{os.getpid()}")
            shutil.copy2(action["source_path"], temp_path)
            if file_sha256(temp_path) != action["source_sha256"]:
                temp_path.unlink(missing_ok=True)
                raise ValueError(f"Copied file hash mismatch: {target_path}")
            os.replace(temp_path, target_path)
            copied.append(target_path)
        return copied
    except Exception:
        for path in copied:
            path.unlink(missing_ok=True)
        raise


def insert_image_rows(container, database, actions):
    inserts = [action for action in actions if action["insert_required"]]
    if not inserts:
        return 0

    values = []
    for action in inserts:
        values.append(
            "(NOW(6), NOW(6), "
            f"{sql_string(action['image_url'])}, {1 if action['is_primary'] else 0}, "
            f"{action['asset_id']})"
        )
    sql = (
        "START TRANSACTION;\n"
        "INSERT INTO asset_images "
        "(created_at, updated_at, image_url, is_primary, asset_id) VALUES\n"
        + ",\n".join(values)
        + ";\nCOMMIT;"
    )
    run_mysql(container, database, sql)
    return len(inserts)


def write_manifest(actions):
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    manifest_path = PROJECT_ROOT / "outputs" / f"asset_image_migration_{timestamp}.csv"
    manifest_path.parent.mkdir(parents=True, exist_ok=True)
    with manifest_path.open("w", newline="", encoding="utf-8-sig") as output:
        writer = csv.DictWriter(
            output,
            fieldnames=[
                "asset_no",
                "source_path",
                "target_path",
                "image_url",
                "is_primary",
                "copy_status",
                "db_status",
            ],
        )
        writer.writeheader()
        for action in actions:
            writer.writerow(
                {
                    "asset_no": action["asset_no"],
                    "source_path": action["source_path"],
                    "target_path": action["target_path"],
                    "image_url": action["image_url"],
                    "is_primary": action["is_primary"],
                    "copy_status": "copied" if action["copy_required"] else "existing",
                    "db_status": "inserted" if action["insert_required"] else "existing",
                }
            )
    return manifest_path


def summarize(actions, apply_mode):
    return {
        "mode": "apply" if apply_mode else "dry-run",
        "mapping_rows": len(actions),
        "unique_assets": len({action["asset_no"] for action in actions}),
        "files_to_copy": sum(action["copy_required"] for action in actions),
        "db_rows_to_insert": sum(action["insert_required"] for action in actions),
        "primary_images_to_assign": sum(action["is_primary"] for action in actions),
        "total_bytes_to_copy": sum(
            action["source_path"].stat().st_size
            for action in actions
            if action["copy_required"]
        ),
    }


def main():
    args = parse_args()
    workbook_path = args.workbook.resolve()
    upload_dir = args.upload_dir.resolve()

    records = load_mapping(workbook_path)
    assets, existing_images = load_database_state(args.container, args.database)
    validate_database_matches(records, assets)
    actions = prepare_actions(records, upload_dir, assets, existing_images)
    print(json.dumps(summarize(actions, args.apply), indent=2))

    if not args.apply:
        print("Dry-run complete. Re-run with --apply after reviewing the summary.")
        return 0

    copied_files = copy_files(actions, upload_dir)
    try:
        inserted_count = insert_image_rows(args.container, args.database, actions)
    except Exception:
        for path in copied_files:
            path.unlink(missing_ok=True)
        raise

    manifest_path = write_manifest(actions)
    print(f"Copied files: {len(copied_files)}")
    print(f"Inserted DB rows: {inserted_count}")
    print(f"Manifest: {manifest_path}")
    return 0


if __name__ == "__main__":
    try:
        sys.exit(main())
    except Exception as error:
        print(f"ERROR: {error}", file=sys.stderr)
        sys.exit(1)
