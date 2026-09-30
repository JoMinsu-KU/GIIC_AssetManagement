import openpyxl
import os
import datetime
import math

# 설정
EXCEL_FILE = r"C:\Users\choiLee\GIIC_Asset\자 산 등 재 목 록.xlsx"
OUTPUT_SQL = "DB/init/10_asset_migration.sql"
REGISTERED_BY = "2025009"

def escape_sql_string(val):
    if val is None:
        return "NULL"
    
    val_str = str(val).strip()
    if not val_str:
        return "NULL"
        
    # 작은따옴표 이스케이프 처리
    escaped = val_str.replace("'", "''")
    return f"'{escaped}'"

def format_date(val):
    if val is None:
        return "NULL"
        
    if isinstance(val, datetime.datetime):
        return f"'{val.strftime('%Y-%m-%d')}'"
        
    val_str = str(val).strip()
    if not val_str:
        return "NULL"
        
    # 문자열 파싱 시도 (예: 2025.12.30, 2025-12-30)
    cleaned = val_str.replace(".", "-").replace("/", "-").strip()
    
    # 시간 정보가 있는 경우 제거
    if " " in cleaned:
        cleaned = cleaned.split(" ")[0]
        
    return f"'{cleaned}'"

def format_number(val):
    if val is None:
        return "0"
        
    if isinstance(val, (int, float)):
        if math.isnan(val):
            return "0"
        return str(int(val))
        
    val_str = str(val).strip().replace(",", "").replace("원", "")
    if not val_str:
        return "0"
        
    try:
        return str(int(float(val_str)))
    except ValueError:
        return "0"

def convert_status(status_raw):
    if status_raw is None:
        return "NORMAL"
        
    status_str = str(status_raw).strip()
    
    if status_str == "정상":
        return "NORMAL"
    elif status_str == "불용":
        return "DISPOSED"
    elif status_str == "고장":
        return "ABNORMAL"
    elif status_str == "수리중":
        return "REPAIRING"
    else:
        return "NORMAL"

def main():
    print(f"Excel 파일 읽는 중... ({EXCEL_FILE})")
    try:
        wb = openpyxl.load_workbook(EXCEL_FILE, data_only=True)
        sheet = wb.active
    except Exception as e:
        print(f"오류: Excel 파일을 열 수 없습니다. ({e})")
        return

    # 헤더 인덱스 찾기
    headers = {}
    header_row_idx = 1
    
    for row_idx, row in enumerate(sheet.iter_rows(min_row=1, max_row=20, values_only=True), 1):
        if not row:
            continue
            
        row_str = [str(cell).strip() if cell else "" for cell in row]
        if any("자산번호" in cell for cell in row_str):
            header_row_idx = row_idx
            for col_idx, cell_value in enumerate(row_str):
                if cell_value:
                    # 줄바꿈, 공백 제거하여 키 생성
                    clean_key = cell_value.replace("\n", "").replace(" ", "")
                    headers[clean_key] = col_idx
            break

    print(f"찾은 헤더 정보: {headers}")

    # 필수 컬럼 검증 (공백 제거된 키로 확인)
    required_cols = ['자산번호', '물품명']
    for col in required_cols:
        if col not in headers:
            print(f"오류: 필수 컬럼 '{col}'을 찾을 수 없습니다.")
            return

    sql_lines = []
    sql_lines.append("START TRANSACTION;\n")

    # DB 컬럼과 매핑되는 엑셀 컬럼 이름
    col_mapping = {
        'asset_no': '자산번호',
        'name': '물품명',
        'model': '모델명',
        'spec': '규격',
        'value': '자산금액',
        'purchase_date': '취득일자',
        'managing_dept': '관리부서',
        'location': '비치장소',
        'status': '상태'
    }

    insert_template = (
        "INSERT INTO assets (asset_no, name, model, spec, value, purchase_date, managing_dept, location, status, registered_by, actual_user_name, actual_user_department, actual_user_contact, created_at, updated_at)\n"
        "VALUES ({asset_no}, {name}, {model}, {spec}, {value}, {purchase_date}, {managing_dept}, {location}, {status}, '{registered_by}', NULL, NULL, NULL, NOW(), NOW());"
    )

    success_count = 0
    # 데이터 행 파싱
    for row in sheet.iter_rows(min_row=header_row_idx + 1, values_only=True):
        asset_no_idx = headers.get('자산번호')
        
        # 자산번호가 없는 행(빈 행)은 건너뜀
        if asset_no_idx is None or asset_no_idx >= len(row) or row[asset_no_idx] is None:
            continue
            
        asset_no_raw = str(row[asset_no_idx]).strip()
        if not asset_no_raw:
            continue

        # 각 필드 데이터 추출
        row_data = {}
        for db_col, excel_col in col_mapping.items():
            idx = headers.get(excel_col)
            val = row[idx] if idx is not None and idx < len(row) else None
            row_data[db_col] = val

        # 데이터 변환 (NULL 처리, 타입 변환, 이스케이프)
        sql_data = {
            'asset_no': escape_sql_string(row_data['asset_no']),
            'name': escape_sql_string(row_data['name']),
            'model': escape_sql_string(row_data['model']),
            'spec': escape_sql_string(row_data['spec']),
            'value': format_number(row_data['value']),
            'purchase_date': format_date(row_data['purchase_date']),
            'managing_dept': escape_sql_string(row_data['managing_dept']),
            'location': escape_sql_string(row_data['location']),
            'status': f"'{convert_status(row_data['status'])}'",
            'registered_by': REGISTERED_BY
        }

        # SQL 생성
        sql = insert_template.format(**sql_data)
        sql_lines.append(sql)
        success_count += 1

    sql_lines.append("\nCOMMIT;")

    # 출력 디렉토리 확인
    os.makedirs(os.path.dirname(OUTPUT_SQL), exist_ok=True)
    
    # SQL 파일로 저장
    with open(OUTPUT_SQL, 'w', encoding='utf-8') as f:
        f.write("\n".join(sql_lines))

    print(f"\n성공적으로 변환 완료!")
    print(f"처리된 자산 수: {success_count}건")
    print(f"SQL 파일 저장: {OUTPUT_SQL}")

if __name__ == "__main__":
    main()
